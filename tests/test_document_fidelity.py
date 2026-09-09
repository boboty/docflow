from __future__ import annotations

import json
import re
import shutil
from pathlib import Path

import openpyxl
import pytest
import yaml
from PIL import Image as PILImage

from docflow.application.generation import PASS, generate_batch
from docflow.catalog.loader import load_catalog
import docflow.catalog.mutation as mutation
from docflow.catalog.mutation import CatalogMutationError, import_seal
from docflow.domain.document import DocumentType
from docflow.templates.registry import TemplateRegistry


def _png(path: Path, color: tuple[int, int, int, int]) -> Path:
    PILImage.new("RGBA", (120, 120), color).save(path, "PNG")
    return path


def _catalog(root: Path, seller_name: str) -> Path:
    root.mkdir(parents=True)
    (root / "organizations.yaml").write_text(
        yaml.safe_dump({"organizations": {"linyi_yier": {
            "name": seller_name, "roles": ["supplier"], "aliases": ["临沂亦尔"]
        }}}, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )
    return root


def _participant_catalog(root: Path, record: dict, sealed: set[str]) -> Path:
    root.mkdir(parents=True)
    organizations = {
        "guangzhou_yier": {"name": record["buyer"], "roles": ["buyer"], "aliases": ["广州亦尔"]},
        "linyi_yier": {"name": record["seller"], "roles": ["supplier"], "aliases": ["临沂亦尔"]},
        "zhongyi": {"name": record["ship_to"]["company"], "roles": ["ship_to"], "aliases": ["众壹"]},
    }
    (root / "organizations.yaml").write_text(
        yaml.safe_dump({"organizations": organizations}, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    colors = {
        "guangzhou_yier": (220, 0, 0, 180),
        "linyi_yier": (0, 0, 220, 180),
        "zhongyi": (0, 160, 0, 180),
    }
    for organization in sealed:
        import_seal(root, organization, _png(root.parent / f"{organization}.source.png", colors[organization]), False)
    return root


def _batch(path: Path, record: dict) -> Path:
    path.write_text(json.dumps([record], ensure_ascii=False), encoding="utf-8")
    return path


# Fields the template's own PrintProfile now authoritatively controls (see
# docflow.renderers.xlsx._apply_print_profile) - the renderer no longer
# preserves these from the source template verbatim.
_PROFILE_CONTROLLED_KEYS = {"paperSize", "orientation", "scale", "fitToWidth", "fitToHeight", "fitToPage", "print_area"}


def _print_signature(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path)
    ws = workbook[workbook.sheetnames[0]]
    setup = ws.page_setup
    margins = ws.page_margins
    options = ws.print_options
    return {
        "paperSize": setup.paperSize,
        "orientation": setup.orientation,
        "scale": setup.scale,
        "fitToWidth": setup.fitToWidth,
        "fitToHeight": setup.fitToHeight,
        "fitToPage": ws.sheet_properties.pageSetUpPr.fitToPage,
        "margins": (margins.left, margins.right, margins.top, margins.bottom, margins.header, margins.footer),
        "options": (options.horizontalCentered, options.verticalCentered, options.gridLines, options.headings),
        "print_area": str(ws.print_area),
        "print_title_rows": ws.print_title_rows,
        "print_title_cols": ws.print_title_cols,
        "row_breaks": tuple(brk.id for brk in ws.row_breaks.brk),
        "col_breaks": tuple(brk.id for brk in ws.col_breaks.brk),
        "header": ws.oddHeader.center.text,
        "footer": ws.oddFooter.right.text,
    }


def _print_area_range(raw: str) -> str:
    """openpyxl's print_area getter returns a fully-qualified absolute
    reference (e.g. "'采购合同'!$A$1:$I$49"); strip the sheet-name prefix and
    $ anchors down to the bare range a PrintProfile declares (e.g. "A1:I49").
    """
    match = re.search(r"\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", raw)
    assert match, f"unrecognized print_area format: {raw!r}"
    return f"{match.group(1)}{match.group(2)}:{match.group(3)}{match.group(4)}"


def _preserved_print_signature(path: Path) -> dict:
    """The subset of `_print_signature()` the renderer must still carry over
    from the source template untouched (margins, headers/footers, breaks,
    ...) - everything except the scale-critical fields the template's own
    PrintProfile now authors.
    """
    return {key: value for key, value in _print_signature(path).items() if key not in _PROFILE_CONTROLLED_KEYS}


def test_import_seal_managed_copy_survives_source_delete_and_workspace_move(tmp_path: Path):
    workspace = tmp_path / "workspace"
    root = _catalog(workspace / ".docflow" / "catalog", "临沂亦尔科技有限公司")
    source = _png(tmp_path / "external.png", (220, 0, 0, 180))

    import_seal(root, "linyi_yier", source, replace=False)
    source.unlink()
    org = load_catalog(root).organizations["linyi_yier"]
    assert org.seal is not None and org.seal.path == "../assets/seals/linyi_yier.png"
    assert (root / org.seal.path).resolve().is_file()

    moved = tmp_path / "moved-workspace"
    shutil.copytree(workspace, moved)
    moved_root = moved / ".docflow" / "catalog"
    moved_org = load_catalog(moved_root).organizations["linyi_yier"]
    assert (moved_root / moved_org.seal.path).resolve().is_file()


def test_import_seal_replace_and_invalid_inputs_are_atomic(tmp_path: Path):
    root = _catalog(tmp_path / "workspace" / ".docflow" / "catalog", "临沂亦尔科技有限公司")
    first = _png(tmp_path / "first.png", (200, 0, 0, 255))
    second = _png(tmp_path / "second.png", (0, 0, 200, 255))
    import_seal(root, "linyi_yier", first, replace=False)
    managed = root.parent / "assets" / "seals" / "linyi_yier.png"
    first_bytes = managed.read_bytes()

    with pytest.raises(CatalogMutationError) as existing:
        import_seal(root, "linyi_yier", second, replace=False)
    assert existing.value.code == "SEAL_ALREADY_EXISTS"
    assert managed.read_bytes() == first_bytes

    import_seal(root, "linyi_yier", second, replace=True)
    second_bytes = managed.read_bytes()
    assert second_bytes != first_bytes

    catalog_before = (root / "organizations.yaml").read_bytes()
    bad = tmp_path / "bad.png"
    bad.write_bytes(b"not a png")
    with pytest.raises(CatalogMutationError) as invalid:
        import_seal(root, "linyi_yier", bad, replace=True)
    assert invalid.value.code == "SEAL_SOURCE_INVALID"
    assert managed.read_bytes() == second_bytes
    assert (root / "organizations.yaml").read_bytes() == catalog_before

    with pytest.raises(CatalogMutationError) as missing:
        import_seal(root, "linyi_yier", tmp_path / "missing.png", replace=True)
    assert missing.value.code == "SEAL_SOURCE_MISSING"
    assert managed.read_bytes() == second_bytes


def test_import_seal_commit_failure_rolls_back_binary_and_catalog(tmp_path: Path, monkeypatch):
    root = _catalog(tmp_path / "workspace" / ".docflow" / "catalog", "临沂亦尔科技有限公司")
    first = _png(tmp_path / "first.png", (200, 0, 0, 255))
    second = _png(tmp_path / "second.png", (0, 0, 200, 255))
    import_seal(root, "linyi_yier", first, replace=False)
    managed = root.parent / "assets" / "seals" / "linyi_yier.png"
    managed_before = managed.read_bytes()
    catalog_before = (root / "organizations.yaml").read_bytes()

    def fail_write(*args, **kwargs):
        raise OSError("simulated catalog write failure")

    monkeypatch.setattr(mutation, "_write_yaml_atomic", fail_write)
    with pytest.raises(CatalogMutationError) as failed:
        import_seal(root, "linyi_yier", second, replace=True)
    assert failed.value.code == "SEAL_IMPORT_FAILED"
    assert managed.read_bytes() == managed_before
    assert (root / "organizations.yaml").read_bytes() == catalog_before
    assert not list(managed.parent.glob(".*.tmp"))
    assert not list(managed.parent.glob(".*.bak"))


def test_import_seal_crops_transparent_padding_and_centers_on_square_canvas(tmp_path: Path):
    """A source PNG's incidental transparent padding must not survive into
    the managed asset: a mapping's printed_diameter_mm describes the seal's
    visible outer size, so any leftover padding baked into the canvas would
    make the renderer's printed-size math shrink the visible seal below the
    declared diameter.
    """
    root = _catalog(tmp_path / "workspace" / ".docflow" / "catalog", "示例卖方有限公司")
    source = tmp_path / "padded.png"
    canvas = PILImage.new("RGBA", (200, 200), (0, 0, 0, 0))
    seal = PILImage.new("RGBA", (60, 40), (255, 0, 0, 255))
    canvas.paste(seal, (70, 90), seal)
    canvas.save(source, "PNG")

    import_seal(root, "linyi_yier", source, replace=False)

    managed = root.parent / "assets" / "seals" / "linyi_yier.png"
    with PILImage.open(managed) as result:
        result = result.convert("RGBA")
        assert result.size == (60, 60)
        assert result.getchannel("A").getbbox() == (0, 10, 60, 50)
        assert result.getpixel((30, 30)) == (255, 0, 0, 255)
        assert result.getpixel((0, 0))[3] == 0

    with PILImage.open(source) as original:
        assert original.size == (200, 200)


def test_import_seal_normalization_preserves_semi_transparent_pixel_values(tmp_path: Path):
    """Regression: normalization must copy cropped pixels verbatim, not
    paste them using their own alpha band as the paste mask. Image.paste()
    with a mask blends source and destination - including the alpha channel
    - for any partially-transparent pixel, which double-applies alpha (a
    128-alpha pixel pasted onto a fully-transparent background would come
    out as (R/2, G/2, B/2, 64) instead of the original (R, G, B, 128)).
    Every non-transparent pixel's RGBA value must be bit-for-bit identical
    before and after normalization.
    """
    root = _catalog(tmp_path / "workspace" / ".docflow" / "catalog", "示例卖方有限公司")
    source = tmp_path / "semi_transparent.png"
    original_pixel = (200, 150, 50, 128)
    canvas = PILImage.new("RGBA", (100, 100), (0, 0, 0, 0))
    translucent_seal = PILImage.new("RGBA", (20, 20), original_pixel)
    canvas.paste(translucent_seal, (40, 40))
    canvas.save(source, "PNG")

    import_seal(root, "linyi_yier", source, replace=False)

    managed = root.parent / "assets" / "seals" / "linyi_yier.png"
    with PILImage.open(managed) as result:
        result = result.convert("RGBA")
        assert result.size == (20, 20)
        for point in ((0, 0), (10, 10), (19, 19)):
            assert result.getpixel(point) == original_pixel


def _generated(
    tmp_path: Path,
    sample_batch_dict: dict,
    contract_template_path: Path,
    delivery_template_path: Path,
    sealed: set[str],
):
    root = _participant_catalog(tmp_path / "workspace" / ".docflow" / "catalog", sample_batch_dict, sealed)
    output = tmp_path / "output"
    summary = generate_batch(
        _batch(tmp_path / "batch.json", sample_batch_dict),
        contract_template_path, delivery_template_path, output,
        catalog_root=root,
    )
    return summary, output


def _entry(summary, document_type: str):
    return next(entry for entry in summary.entries if entry.document_type == document_type)


@pytest.mark.parametrize(
    ("document_type", "sealed", "inserted", "skipped"),
    [
        ("procurement.contract.v1", {"guangzhou_yier"}, {"buyer_seal"}, {"seller_seal"}),
        ("procurement.contract.v1", {"linyi_yier"}, {"seller_seal"}, {"buyer_seal"}),
        ("procurement.contract.v1", {"guangzhou_yier", "linyi_yier"}, {"buyer_seal", "seller_seal"}, set()),
        ("delivery.note.v1", {"linyi_yier"}, {"seller_seal"}, {"ship_to_seal"}),
        ("delivery.note.v1", {"zhongyi"}, {"ship_to_seal"}, {"seller_seal"}),
        ("delivery.note.v1", {"linyi_yier", "zhongyi"}, {"seller_seal", "ship_to_seal"}, set()),
    ],
)
def test_participant_seals_are_resolved_and_inserted_independently(
    tmp_path: Path,
    sample_batch_dict: dict,
    contract_template_path: Path,
    delivery_template_path: Path,
    document_type: str,
    sealed: set[str],
    inserted: set[str],
    skipped: set[str],
):
    summary, output = _generated(
        tmp_path, sample_batch_dict, contract_template_path, delivery_template_path, sealed,
    )

    assert summary.failed_documents == 0
    assert all(entry.validation_status == PASS for entry in summary.entries)
    entry = _entry(summary, document_type)
    assert len(entry.enhancements) == 2
    for image_id in inserted:
        assert f"IMAGE_INSERTED {image_id}" in entry.enhancements
    for image_id in skipped:
        matching = [note for note in entry.enhancements if note.startswith(f"IMAGE_SKIPPED {image_id}:")]
        assert len(matching) == 1
        assert "asset not configured" not in matching[0]
    assert len(openpyxl.load_workbook(output / entry.output_file).active._images) == len(inserted)


@pytest.mark.parametrize("resolution", ["missing", "ambiguous"])
def test_missing_or_ambiguous_participant_only_skips_that_slot(
    tmp_path: Path, sample_batch_dict: dict, contract_template_path: Path, delivery_template_path: Path,
    resolution: str,
):
    root = _participant_catalog(
        tmp_path / "workspace" / ".docflow" / "catalog",
        sample_batch_dict,
        {"guangzhou_yier", "linyi_yier"},
    )
    raw = yaml.safe_load((root / "organizations.yaml").read_text(encoding="utf-8"))
    if resolution == "missing":
        raw["organizations"].pop("guangzhou_yier")
    else:
        duplicate = dict(raw["organizations"]["guangzhou_yier"])
        duplicate["seal"] = raw["organizations"]["linyi_yier"]["seal"]
        raw["organizations"]["guangzhou_yier_duplicate"] = duplicate
    (root / "organizations.yaml").write_text(
        yaml.safe_dump(raw, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    output = tmp_path / "output"
    summary = generate_batch(
        _batch(tmp_path / "batch.json", sample_batch_dict),
        contract_template_path, delivery_template_path, output,
        catalog_root=root,
    )
    assert summary.failed_documents == 0
    contract = _entry(summary, "procurement.contract.v1")
    assert f"IMAGE_SKIPPED buyer_seal: buyer organization {resolution}" in contract.enhancements
    assert "IMAGE_INSERTED seller_seal" in contract.enhancements
    assert len(openpyxl.load_workbook(output / contract.output_file).active._images) == 1
    assert all("asset not configured" not in note for note in contract.enhancements)


def test_multi_seal_generation_applies_print_profile_and_expected_anchors(
    tmp_path: Path, sample_batch_dict: dict, contract_template_path: Path, delivery_template_path: Path,
):
    summary, output = _generated(
        tmp_path, sample_batch_dict, contract_template_path, delivery_template_path,
        {"guangzhou_yier", "linyi_yier", "zhongyi"},
    )
    profiles = {
        "procurement.contract.v1": TemplateRegistry().get(DocumentType.PROCUREMENT_CONTRACT_V1).print_profile,
        "delivery.note.v1": TemplateRegistry().get(DocumentType.DELIVERY_NOTE_V1).print_profile,
    }
    expected = {
        "procurement.contract.v1": ({(0, 45), (5, 45)}, contract_template_path),
        "delivery.note.v1": ({(1, 30), (5, 30)}, delivery_template_path),
    }
    for entry in summary.entries:
        generated = output / entry.output_file
        ws = openpyxl.load_workbook(generated).active
        anchors = {(image.anchor._from.col, image.anchor._from.row) for image in ws._images}
        expected_anchors, template = expected[entry.document_type]
        assert anchors == expected_anchors

        # Printed seal size is fixed at 38mm on paper; since the sheet
        # itself prints at profile.scale_percent, the workbook-embedded
        # image must be inflated by the inverse of that factor so the
        # PRINTED result still measures 38mm (Repair-3 fix: this used to be
        # a bare 38mm workbook size, which shrank on paper along with the
        # rest of the fit-to-page'd sheet).
        profile = profiles[entry.document_type]
        expected_emu = round(38 / (profile.scale_percent / 100.0) * 36_000)
        assert all((image.anchor.ext.cx, image.anchor.ext.cy) == (expected_emu, expected_emu) for image in ws._images)

        # Non-scale settings (margins, headers/footers, breaks, ...) still
        # come straight from the source template, untouched.
        assert _preserved_print_signature(generated) == _preserved_print_signature(template)

        # Scale-critical settings are authored from the template's own
        # calibrated PrintProfile, NOT preserved from the source template
        # (whose synthetic fixture deliberately carries a stale, dynamic
        # fit-to-page setting to prove it is overridden, not honored).
        signature = _print_signature(generated)
        assert signature["paperSize"] == int(ws.PAPERSIZE_A4)
        assert signature["orientation"] == profile.orientation
        assert signature["scale"] == profile.scale_percent
        assert signature["fitToWidth"] is None
        assert signature["fitToHeight"] is None
        assert signature["fitToPage"] is False
        assert _print_area_range(signature["print_area"]) == profile.print_area
