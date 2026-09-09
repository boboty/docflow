"""Managed Template Lifecycle: `docflow catalog import-template`.

Covers the external-source -> workspace-managed-asset transition: a real
template file the user points to is a *source*; once explicitly saved,
the Catalog owns a copy of it under the workspace's own
.docflow/templates/ and never depends on the external location again.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest

from docflow import cli
from docflow.catalog.loader import load_catalog
from docflow.catalog import mutation
from docflow.catalog.mutation import CatalogMutationError, import_template
from docflow.catalog.resolve import resolve_template
from docflow.catalog.root import managed_template_root
from docflow.renderers.xlsx import TemplatePreflightError
from tests.fixtures.synthetic_templates import build_contract_template, build_delivery_template

CONTRACT_DOC_TYPE = "procurement.contract.v1"
DELIVERY_DOC_TYPE = "delivery.note.v1"


def _catalog_root(tmp_path: Path) -> Path:
    return tmp_path / "workspace" / ".docflow" / "catalog"


# 1. fresh import
def test_fresh_import_creates_managed_copy_and_resolvable_entry(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    external_dir = tmp_path / "external"
    external_dir.mkdir()
    source = build_contract_template(external_dir / "采购合同_YE20260516003.xlsx")

    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, source, replace=False)

    managed_file = managed_template_root(catalog_root) / "procurement_contract.xlsx"
    assert managed_file.exists()

    catalog = load_catalog(catalog_root)
    entry = catalog.templates["procurement_contract"]
    assert entry.path == "../templates/procurement_contract.xlsx"  # relative, not the external absolute path
    assert "external" not in entry.path

    result = resolve_template(catalog, "procurement_contract", catalog_root=catalog_root)
    assert result["status"] == "RESOLVED"
    assert Path(result["path"]) == managed_file.resolve()
    assert Path(result["path"]).exists()


# 2. source remains external, and is safely disposable after import
def test_external_source_untouched_and_disposable_after_import(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    external_dir = tmp_path / "external"
    external_dir.mkdir()
    source = build_contract_template(external_dir / "采购合同_YE20260516003.xlsx")
    original_bytes = source.read_bytes()

    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, source, replace=False)

    # Not modified.
    assert source.read_bytes() == original_bytes

    # Not deleted, and now delete it ourselves to prove the Catalog no
    # longer depends on it.
    source.unlink()

    catalog = load_catalog(catalog_root)
    result = resolve_template(catalog, "procurement_contract", catalog_root=catalog_root)
    assert result["status"] == "RESOLVED"
    assert Path(result["path"]).exists()  # the MANAGED copy, independent of the deleted source


# 3. workspace move
def test_workspace_move_keeps_template_resolvable_from_new_location(tmp_path: Path):
    workspace_a = tmp_path / "workspace-a"
    catalog_root_a = workspace_a / ".docflow" / "catalog"
    source = build_contract_template(tmp_path / "external.xlsx")

    import_template(catalog_root_a, "procurement_contract", CONTRACT_DOC_TYPE, source, replace=False)

    workspace_b = tmp_path / "workspace-b"
    shutil.copytree(workspace_a, workspace_b)
    catalog_root_b = workspace_b / ".docflow" / "catalog"

    catalog_b = load_catalog(catalog_root_b)
    result = resolve_template(catalog_b, "procurement_contract", catalog_root=catalog_root_b)
    assert result["status"] == "RESOLVED"
    resolved_path = Path(result["path"])
    assert resolved_path.exists()
    assert str(workspace_b) in str(resolved_path)
    assert str(workspace_a) not in str(resolved_path)


# 4. existing template without --replace
def test_import_without_replace_rejects_existing_key(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    source_1 = build_contract_template(tmp_path / "first.xlsx")
    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, source_1, replace=False)

    managed_file = managed_template_root(catalog_root) / "procurement_contract.xlsx"
    original_managed_bytes = managed_file.read_bytes()
    original_yaml = (catalog_root / "templates.yaml").read_text(encoding="utf-8")

    source_2 = build_contract_template(tmp_path / "second.xlsx")
    with pytest.raises(CatalogMutationError) as exc_info:
        import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, source_2, replace=False)
    assert exc_info.value.code == "TEMPLATE_ALREADY_EXISTS"

    assert managed_file.read_bytes() == original_managed_bytes
    assert (catalog_root / "templates.yaml").read_text(encoding="utf-8") == original_yaml


# 5. replace
def test_import_with_replace_overwrites_managed_copy(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    source_1 = build_contract_template(tmp_path / "first.xlsx")
    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, source_1, replace=False)

    managed_file = managed_template_root(catalog_root) / "procurement_contract.xlsx"
    first_bytes = managed_file.read_bytes()

    source_2 = build_contract_template(tmp_path / "second.xlsx")
    wb = openpyxl.load_workbook(source_2)
    wb["Sheet1"]["C2"] = "采购合同 v2"  # make it distinguishably different
    wb.save(source_2)

    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, source_2, replace=True)

    assert managed_file.read_bytes() != first_bytes
    catalog = load_catalog(catalog_root)
    assert catalog.templates["procurement_contract"].path == "../templates/procurement_contract.xlsx"
    result = resolve_template(catalog, "procurement_contract", catalog_root=catalog_root)
    assert result["status"] == "RESOLVED"


# 6. bad source: missing file, corrupt xlsx, wrong sheet, incompatible mapping
def test_import_rejects_missing_source(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    with pytest.raises(TemplatePreflightError) as exc_info:
        import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, tmp_path / "nope.xlsx", replace=False)
    assert exc_info.value.code == "TEMPLATE_FILE_MISSING"
    assert not catalog_root.exists()
    assert not managed_template_root(catalog_root).exists()


def test_import_rejects_corrupt_xlsx(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    bad_source = tmp_path / "corrupt.xlsx"
    bad_source.write_text("not a zip file", encoding="utf-8")

    with pytest.raises(TemplatePreflightError) as exc_info:
        import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, bad_source, replace=False)
    assert exc_info.value.code == "TEMPLATE_FILE_INVALID"
    assert not managed_template_root(catalog_root).exists()


def test_import_rejects_wrong_sheet(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    wrong_sheet_source = tmp_path / "wrong-sheet.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "NotTheRightSheet"
    wb.save(wrong_sheet_source)

    with pytest.raises(TemplatePreflightError) as exc_info:
        import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, wrong_sheet_source, replace=False)
    assert exc_info.value.code == "TEMPLATE_SHEET_MISSING"
    assert not managed_template_root(catalog_root).exists()


def test_import_rejects_incompatible_mapping(tmp_path: Path):
    """A source with the right sheet name but a header-mapped cell that
    isn't a writable merge anchor - the same preflight docflow generate
    would run must reject it here too, before anything is copied."""
    catalog_root = _catalog_root(tmp_path)
    incompatible_source = tmp_path / "incompatible.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # C34 is a header-mapped cell in the real mapping; merging C33:C34
    # makes C33 the anchor and C34 a non-writable MergedCell.
    ws.merge_cells("C33:C34")
    wb.save(incompatible_source)

    with pytest.raises(TemplatePreflightError) as exc_info:
        import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, incompatible_source, replace=False)
    assert exc_info.value.code == "TEMPLATE_MAPPING_INVALID"
    assert not managed_template_root(catalog_root).exists()


def test_import_rejects_unsupported_document_type(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    source = build_contract_template(tmp_path / "external.xlsx")
    with pytest.raises(CatalogMutationError) as exc_info:
        import_template(catalog_root, "procurement_contract", "not.a.real.type", source, replace=False)
    assert exc_info.value.code == "UNSUPPORTED_DOCUMENT_TYPE"
    assert not managed_template_root(catalog_root).exists()


# 7. existing absolute-path Catalog compatibility
def test_legacy_absolute_path_entry_is_still_valid_and_resolvable(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    catalog_root.mkdir(parents=True)
    external = tmp_path / "external-template.xlsx"
    external.touch()
    (catalog_root / "templates.yaml").write_text(f"""
templates:
  procurement_contract:
    document_type: {CONTRACT_DOC_TYPE}
    path: {external}
""", encoding="utf-8")

    # validate does not reject the whole catalog for this.
    catalog = load_catalog(catalog_root)  # must not raise
    result = resolve_template(catalog, "procurement_contract", catalog_root=catalog_root)
    assert result["status"] == "RESOLVED"
    assert result["path"] == str(external)

    # Explicit re-registration converts it to a managed copy.
    real_source = build_contract_template(tmp_path / "real-source.xlsx")
    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, real_source, replace=True)

    catalog = load_catalog(catalog_root)
    assert catalog.templates["procurement_contract"].path == "../templates/procurement_contract.xlsx"


def test_delivery_note_import_uses_its_own_canonical_filename(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    source = build_delivery_template(tmp_path / "送货单_FBA20260620009.xlsx")
    import_template(catalog_root, "delivery_note", DELIVERY_DOC_TYPE, source, replace=False)

    managed_file = managed_template_root(catalog_root) / "delivery_note.xlsx"
    assert managed_file.exists()
    catalog = load_catalog(catalog_root)
    assert catalog.templates["delivery_note"].path == "../templates/delivery_note.xlsx"


@pytest.mark.parametrize("key", ["../../evil", "foo/bar", "unknown"])
def test_import_rejects_noncanonical_template_key_before_creating_managed_paths(tmp_path: Path, key: str):
    catalog_root = _catalog_root(tmp_path)
    source = build_contract_template(tmp_path / "external.xlsx")

    with pytest.raises(CatalogMutationError) as exc_info:
        import_template(catalog_root, key, CONTRACT_DOC_TYPE, source, replace=False)

    assert exc_info.value.code == "UNSUPPORTED_TEMPLATE_KEY"
    assert not catalog_root.exists()
    assert not managed_template_root(catalog_root).exists()


def test_import_rejects_canonical_key_document_type_mismatch(tmp_path: Path):
    catalog_root = _catalog_root(tmp_path)
    source = build_contract_template(tmp_path / "external.xlsx")

    with pytest.raises(CatalogMutationError) as exc_info:
        import_template(catalog_root, "delivery_note", CONTRACT_DOC_TYPE, source, replace=False)

    assert exc_info.value.code == "TEMPLATE_IDENTITY_MISMATCH"
    assert not managed_template_root(catalog_root).exists()


def _force_yaml_failure(*args, **kwargs):
    raise OSError("forced templates.yaml failure")


def _assert_no_transaction_artifacts(catalog_root: Path) -> None:
    if catalog_root.exists():
        assert not list(catalog_root.rglob("*.tmp"))
        assert not list(catalog_root.rglob("*.bak"))


def test_fresh_import_rolls_back_managed_file_when_yaml_write_fails(tmp_path: Path, monkeypatch, capsys):
    catalog_root = _catalog_root(tmp_path)
    source = build_contract_template(tmp_path / "external.xlsx")
    catalog_path = catalog_root / "templates.yaml"
    catalog_root.mkdir(parents=True)
    original_catalog = b"templates: {}\n"
    catalog_path.write_bytes(original_catalog)
    monkeypatch.setattr(mutation, "_write_yaml_atomic", _force_yaml_failure)

    result = cli.main([
        "catalog", "import-template", "--catalog-root", str(catalog_root),
        "--key", "procurement_contract", "--document-type", CONTRACT_DOC_TYPE,
        "--source", str(source),
    ])

    assert result == 2
    assert "TEMPLATE IMPORT REJECTED" in capsys.readouterr().err
    assert not (managed_template_root(catalog_root) / "procurement_contract.xlsx").exists()
    assert catalog_path.read_bytes() == original_catalog
    _assert_no_transaction_artifacts(catalog_root)


def test_replace_rolls_back_old_managed_bytes_when_yaml_write_fails(tmp_path: Path, monkeypatch, capsys):
    catalog_root = _catalog_root(tmp_path)
    first_source = build_contract_template(tmp_path / "first.xlsx")
    import_template(catalog_root, "procurement_contract", CONTRACT_DOC_TYPE, first_source, replace=False)
    managed_file = managed_template_root(catalog_root) / "procurement_contract.xlsx"
    catalog_path = catalog_root / "templates.yaml"
    original_managed = managed_file.read_bytes()
    original_catalog = catalog_path.read_bytes()

    second_source = build_contract_template(tmp_path / "second.xlsx")
    workbook = openpyxl.load_workbook(second_source)
    workbook["Sheet1"]["C2"] = "replacement"
    workbook.save(second_source)
    assert second_source.read_bytes() != original_managed
    monkeypatch.setattr(mutation, "_write_yaml_atomic", _force_yaml_failure)

    result = cli.main([
        "catalog", "import-template", "--catalog-root", str(catalog_root),
        "--key", "procurement_contract", "--document-type", CONTRACT_DOC_TYPE,
        "--source", str(second_source), "--replace",
    ])

    assert result == 2
    assert "TEMPLATE IMPORT REJECTED" in capsys.readouterr().err
    assert managed_file.read_bytes() == original_managed
    assert catalog_path.read_bytes() == original_catalog
    _assert_no_transaction_artifacts(catalog_root)
