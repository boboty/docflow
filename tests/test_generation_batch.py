import copy
import json
from pathlib import Path

import openpyxl

from docflow.application.generation import FAILED, PASS, generate_batch


def _write_batch(tmp_path: Path, records: list[dict], name: str = "batch.json") -> Path:
    path = tmp_path / name
    path.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
    return path


def test_same_fact_pack_generates_both_documents(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    batch_path = _write_batch(tmp_path, [sample_batch_dict])
    output_dir = tmp_path / "output"

    summary = generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)

    assert summary.total_records == 1
    assert summary.passed_documents == 2
    assert summary.failed_documents == 0

    ref = sample_batch_dict["business_reference"]
    assert (output_dir / ref / "procurement-contract.xlsx").exists()
    assert (output_dir / ref / "delivery-note.xlsx").exists()


def test_one_failing_record_does_not_abort_batch(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    broken = copy.deepcopy(sample_batch_dict)
    broken["business_reference"] = "BR-BROKEN"
    del broken["buyer"]  # missing required field -> should FAIL, not raise

    good = copy.deepcopy(sample_batch_dict)
    good["business_reference"] = "BR-GOOD"

    batch_path = _write_batch(tmp_path, [broken, good])
    output_dir = tmp_path / "output"

    summary = generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)

    assert summary.total_records == 2
    assert summary.passed_documents == 2  # only the good record's 2 documents
    assert summary.failed_documents == 2  # the broken record's 2 documents

    good_entries = [e for e in summary.entries if e.business_reference == "BR-GOOD"]
    broken_entries = [e for e in summary.entries if e.business_reference == "BR-BROKEN"]
    assert all(e.validation_status == PASS for e in good_entries)
    assert all(e.validation_status == FAILED for e in broken_entries)
    assert all(e.issues for e in broken_entries)
    assert (output_dir / "BR-GOOD" / "procurement-contract.xlsx").exists()
    assert not (output_dir / "BR-BROKEN").exists()


def test_manifest_is_written_and_structurally_correct(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    batch_path = _write_batch(tmp_path, [sample_batch_dict])
    output_dir = tmp_path / "output"

    generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)

    manifest_path = output_dir / "manifest.json"
    assert manifest_path.exists()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

    assert len(manifest) == 2
    required_fields = {
        "business_reference", "document_type", "template_id", "template_version",
        "output_file", "validation_status", "issues", "source_snapshot_hash",
    }
    for entry in manifest:
        assert required_fields <= entry.keys()
        assert entry["validation_status"] == "PASS"
        assert (output_dir / entry["output_file"]).exists()


def test_same_input_produces_stable_business_content(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    batch_path = _write_batch(tmp_path, [sample_batch_dict])

    summary_a = generate_batch(batch_path, contract_template_path, delivery_template_path, tmp_path / "out-a")
    summary_b = generate_batch(batch_path, contract_template_path, delivery_template_path, tmp_path / "out-b")

    manifest_a = json.loads((tmp_path / "out-a" / "manifest.json").read_text(encoding="utf-8"))
    manifest_b = json.loads((tmp_path / "out-b" / "manifest.json").read_text(encoding="utf-8"))
    assert manifest_a == manifest_b

    wb_a = openpyxl.load_workbook(tmp_path / "out-a" / summary_a.entries[0].output_file)
    wb_b = openpyxl.load_workbook(tmp_path / "out-b" / summary_b.entries[0].output_file)
    ws_a = wb_a[wb_a.sheetnames[0]]
    ws_b = wb_b[wb_b.sheetnames[0]]
    assert ws_a["A3"].value == ws_b["A3"].value
    assert ws_a["I9"].value == ws_b["I9"].value


def test_contract_and_delivery_files_agree_on_amount_within_source_tolerance(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    """P0 regression: gross_unit_price*quantity is allowed to diverge from
    gross_amount within the documented tolerance. Both generated FILES
    (not just the in-memory projections) must still show the same,
    authoritative amount for that line - not one computed by multiplying
    D*F in the delivery note.
    """
    record = copy.deepcopy(sample_batch_dict)
    record["business_reference"] = "BR-TOLERANCE"
    record["items"] = [{
        "sku": "SKU-042", "product_name": "测试产品", "specification": "型号SKU-042",
        "quantity": 42, "unit": "件", "gross_unit_price": "49", "gross_amount": "2058.10",
        "tax_rate": "0.13",
    }]

    batch_path = _write_batch(tmp_path, [record])
    output_dir = tmp_path / "output"

    summary = generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)
    assert summary.failed_documents == 0, [(e.document_type, e.issues) for e in summary.entries]

    contract_wb = openpyxl.load_workbook(output_dir / "BR-TOLERANCE" / "procurement-contract.xlsx")
    delivery_wb = openpyxl.load_workbook(output_dir / "BR-TOLERANCE" / "delivery-note.xlsx")
    contract_gross = contract_wb["采购合同"]["I9"].value
    delivery_gross = delivery_wb["送货单"]["G7"].value

    assert contract_gross == 2058.10
    assert delivery_gross == 2058.10
    assert contract_gross == delivery_gross
    assert delivery_gross != 42 * 49  # the recomputed-formula value this bug used to show


def test_non_finite_decimal_fails_cleanly_without_crashing_batch(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    good_1 = copy.deepcopy(sample_batch_dict)
    good_1["business_reference"] = "BR-GOOD-1"

    for bad_value in ("NaN", "Infinity", "-Infinity"):
        bad = copy.deepcopy(sample_batch_dict)
        bad["business_reference"] = f"BR-BAD-{bad_value}"
        bad["items"][0]["gross_amount"] = bad_value

        good_2 = copy.deepcopy(sample_batch_dict)
        good_2["business_reference"] = "BR-GOOD-2"

        batch_path = _write_batch(tmp_path, [good_1, bad, good_2], name=f"batch-{bad_value}.json")
        output_dir = tmp_path / f"output-{bad_value}"

        summary = generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)

        assert summary.total_records == 3
        assert summary.passed_documents == 4  # good_1 + good_2, 2 documents each
        assert summary.failed_documents == 2  # bad record only
        bad_entries = [e for e in summary.entries if e.business_reference == bad["business_reference"]]
        assert all(e.validation_status == FAILED for e in bad_entries)
        assert all(
            "NON_FINITE_DECIMAL" in issue or "NON_FINITE_VALUE" in issue
            for e in bad_entries for issue in e.issues
        )
        assert (output_dir / "BR-GOOD-1" / "procurement-contract.xlsx").exists()
        assert (output_dir / "BR-GOOD-2" / "procurement-contract.xlsx").exists()


def test_non_dict_record_fails_cleanly_without_crashing_batch(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    # A batch element that isn't even a JSON object (e.g. a stray string)
    # must not raise an unhandled exception out of generate_batch.
    batch_path = _write_batch(tmp_path, ["not-a-record", sample_batch_dict])
    output_dir = tmp_path / "output"

    summary = generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)

    assert summary.total_records == 2
    assert summary.passed_documents == 2
    assert summary.failed_documents == 2
    failed = [e for e in summary.entries if e.validation_status == FAILED]
    assert all("INVALID_RECORD_SHAPE" in issue for e in failed for issue in e.issues)


def test_malformed_extra_field_fails_cleanly_without_crashing_batch(
    sample_batch_dict, contract_template_path, delivery_template_path, tmp_path
):
    broken = copy.deepcopy(sample_batch_dict)
    broken["business_reference"] = "BR-BAD-EXTRA"
    broken["extra"] = ["not", "an", "object"]  # must not raise dict()/TypeError

    batch_path = _write_batch(tmp_path, [broken])
    output_dir = tmp_path / "output"

    summary = generate_batch(batch_path, contract_template_path, delivery_template_path, output_dir)

    assert summary.total_records == 1
    assert summary.failed_documents == 2
    assert all(e.validation_status == FAILED for e in summary.entries)
    assert all("INVALID_RECORD_SHAPE" in issue for e in summary.entries for issue in e.issues)
