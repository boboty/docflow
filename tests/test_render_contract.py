from dataclasses import replace
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from docflow.domain.document import DocumentType, build_projection
from docflow.renderers.xlsx import TEMPLATE_ITEM_CAPACITY_EXCEEDED, TemplateRenderError, render
from docflow.templates.registry import TemplateRegistry
from tests.conftest import CONTRACT_CAPACITY, make_fact_pack


def _definition():
    return TemplateRegistry().get(DocumentType.PROCUREMENT_CONTRACT_V1)


def test_single_contract_generation(contract_template_path: Path, tmp_path: Path):
    pack = make_fact_pack(item_count=2)
    projection = build_projection(pack, DocumentType.PROCUREMENT_CONTRACT_V1)
    output_path = tmp_path / "out" / "procurement-contract.xlsx"

    render(_definition(), contract_template_path, projection, output_path)

    assert output_path.exists()
    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]
    assert ws["C3"].value == f"购方：{pack.buyer}"
    assert ws["I3"].value == f"合同编号：{pack.contract_no}"
    assert ws["C4"].value == f"销方：{pack.seller}"
    assert ws["D21"].value == projection.amount_in_words

    first_item = projection.items[0]
    assert ws["C8"].value == first_item.sku
    assert ws["D8"].value == first_item.product_name
    assert ws["F8"].value == float(first_item.quantity)
    assert ws["K8"].value == float(first_item.gross_amount)


def test_below_capacity_clears_leftover_sample_rows(contract_template_path: Path, tmp_path: Path):
    pack = make_fact_pack(item_count=2)
    projection = build_projection(pack, DocumentType.PROCUREMENT_CONTRACT_V1)
    output_path = tmp_path / "out.xlsx"

    render(_definition(), contract_template_path, projection, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]
    # Row 10 is the first leftover row (items filled rows 8-9) and the
    # synthetic template pre-seeds it with sample data that must be cleared.
    for column in "CDEFGHIJK":
        assert ws[f"{column}10"].value is None
    for column in "CDEFGHIJK":
        assert ws[f"{column}19"].value is None


def test_at_capacity_exact_fills_all_rows(contract_template_path: Path, tmp_path: Path):
    pack = make_fact_pack(item_count=CONTRACT_CAPACITY)
    projection = build_projection(pack, DocumentType.PROCUREMENT_CONTRACT_V1)
    output_path = tmp_path / "out.xlsx"

    render(_definition(), contract_template_path, projection, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]
    last_item = projection.items[-1]
    assert ws["C19"].value == last_item.sku
    assert ws["C8"].value == projection.items[0].sku


def test_body_text_overwrites_previous_transactions_delivery_date(contract_template_path: Path, tmp_path: Path):
    """C29 regression: the synthetic template's row 29 pre-seeds a stale
    delivery-deadline clause ("6月9日前...") mirroring a real transaction
    that leaked into a real generated contract. This transaction's own
    delivery_date must fully replace it - the old date must not survive
    anywhere in the cell.
    """
    pack = replace(make_fact_pack(item_count=1), delivery_date=date(2026, 9, 9))
    projection = build_projection(pack, DocumentType.PROCUREMENT_CONTRACT_V1)
    output_path = tmp_path / "out.xlsx"

    render(_definition(), contract_template_path, projection, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]
    c29 = ws["C29"].value
    assert "9月9日前" in c29
    assert "6月9日前" not in c29
    assert "6月9日" not in c29


def test_over_capacity_raises_explicit_error(contract_template_path: Path, tmp_path: Path):
    pack = make_fact_pack(item_count=CONTRACT_CAPACITY + 1)
    projection = build_projection(pack, DocumentType.PROCUREMENT_CONTRACT_V1)
    output_path = tmp_path / "out.xlsx"

    with pytest.raises(TemplateRenderError) as exc_info:
        render(_definition(), contract_template_path, projection, output_path)

    assert exc_info.value.code == TEMPLATE_ITEM_CAPACITY_EXCEEDED
    assert not output_path.exists()
