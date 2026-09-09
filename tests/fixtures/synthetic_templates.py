"""Builds synthetic xlsx templates that mirror the *structure* (merged cells,
item table position, subtotal formulas, capacity) of the two real templates
this project targets, using fabricated sample data.

These mirror the ORIGINAL business Excel files (Template Migration, 2026-09)
- sheet named "Sheet1", item table at rows 8-19 / 7-18, native formulas at
the real coordinates - not the earlier OCR/scan reconstruction this project
used before real template files were available.

No real template file, company name, address, phone number or contract
number is ever used here or committed to this repository.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.worksheet.pagebreak import Break

CONTRACT_ITEM_START_ROW = 8
CONTRACT_ITEM_END_ROW = 19  # capacity 12, matches the real template

DELIVERY_ITEM_START_ROW = 7
DELIVERY_ITEM_END_ROW = 18  # capacity 12, matches the real template


def _configure_print(ws, print_area: str, print_title_rows: str, row_break: int) -> None:
    """Deliberately stale, dynamic "Fit to Page" settings - like a real
    template a user has manually zoomed at some point - so tests can prove
    the renderer's own calibrated PrintProfile authoritatively overrides
    them rather than preserving whatever the source file happened to carry.
    """
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.scale = 95
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False
    ws.print_options.gridLines = False
    ws.print_area = print_area
    ws.print_title_rows = print_title_rows
    ws.row_breaks.append(Break(id=row_break))
    ws.col_breaks.append(Break(id=4))
    ws.oddHeader.center.text = "DocFlow 测试页眉"
    ws.oddFooter.right.text = "第 &P 页，共 &N 页"


def build_contract_template(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.merge_cells("C2:K2")
    ws["C2"] = "采购合同"

    ws.merge_cells("C3:H3")
    ws["C3"] = "购方：示例采购方有限公司"
    ws.merge_cells("I3:K3")
    ws["I3"] = "合同编号：SAMPLE-0000"

    ws.merge_cells("C4:K4")
    ws["C4"] = "销方：示例供应商有限公司"

    ws.merge_cells("C5:K5")
    ws["C5"] = "兹经供需双方协商同意，特订立本合同如下。"

    ws["B6"] = "一、"
    ws.merge_cells("C6:K6")
    ws["C6"] = "产品名称、规格、计量单位、数量、单价、金额、交(提)货期限："

    headers = ["SKU", "产品名称", "规格型号", "数量", "单位", "单价（不含税）", "不含税金额", "税金", "价税合计"]
    for col_idx, text in enumerate(headers, start=3):  # starts at column C
        col = openpyxl.utils.get_column_letter(col_idx)
        ws[f"{col}7"] = text

    # Pre-existing sample rows (like a real template shipped with example
    # data) so tests can verify leftover rows get cleared.
    for row in range(CONTRACT_ITEM_START_ROW, CONTRACT_ITEM_END_ROW + 1):
        ws[f"C{row}"] = f"SAMPLE-SKU-{row}"
        ws[f"D{row}"] = "示例产品"
        ws[f"E{row}"] = f"示例型号-{row}"
        ws[f"F{row}"] = 1
        ws[f"G{row}"] = "件"
        ws[f"H{row}"] = 1
        ws[f"I{row}"] = 1
        ws[f"J{row}"] = 0
        ws[f"K{row}"] = 1

    ws.merge_cells("C20:H20")
    ws["C20"] = "小计"
    ws["I20"] = f"=SUM(I{CONTRACT_ITEM_START_ROW}:I{CONTRACT_ITEM_END_ROW})"
    ws["J20"] = f"=SUM(J{CONTRACT_ITEM_START_ROW}:J{CONTRACT_ITEM_END_ROW})"
    ws["K20"] = f"=SUM(K{CONTRACT_ITEM_START_ROW}:K{CONTRACT_ITEM_END_ROW})"

    ws["C21"] = "合并（RMB大写）："
    ws.merge_cells("D21:H21")
    ws["D21"] = "零元整"  # stale native-formula result a real file computes via TEXT/DBNum2
    ws["I21"] = "合并（RMB小写）："
    ws["K21"] = "=K20"

    clause_rows = [
        (22, "二", "质量要求技术标准，乙方对质量负责的条件和期限："),
        (23, None, "销方确保质量应符合购方确认的样品要求和标准。"),
        (24, "三", "交(提)货地点、方式："),
        (25, None, "购方指定地点。"),
        (26, "四", "运输方式和运费负担："),
        (27, None, "货物由销方运到购方指定地点，采购单价包含运费。"),
        (28, "五", "违约责任："),
        # Mirrors the real template bug this fixture exists to regression-test:
        # row 29 is NOT purely static - it embeds a delivery-deadline date left
        # over from whichever transaction last used this physical template
        # file. A correct mapping/renderer must overwrite it with THIS
        # transaction's delivery date, never leave a previous one in place.
        (29, None, "6月9日前分批次完成发货，每逾期1日，乙方按合同总金额的3%支付违约金；逾期超过7日，甲方有权解除合同并要求赔偿损失。"),
        (30, "六", "结算方式及期限："),
        (31, None, "全款，发货前结清。"),
        (32, "七", "争议解决："),
        (33, None, "凡本合同所引起的一切争议，双方应通过友好协商解决。"),
    ]
    for row, label, text in clause_rows:
        if label is not None:
            ws[f"B{row}"] = f"{label}、"
        ws.merge_cells(f"C{row}:K{row}")
        ws[f"C{row}"] = text

    ws.merge_cells("C34:F34")
    ws["C34"] = "购方：示例采购方有限公司"
    ws.merge_cells("G34:K34")
    ws["G34"] = "销方：示例供应商有限公司"
    ws.merge_cells("C35:F35")
    ws["C35"] = "时间："
    ws.merge_cells("G35:K35")
    ws["G35"] = "时间："

    _configure_print(ws, "B1:K35", "1:7", 20)

    wb.save(path)
    return path


def build_delivery_template(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.merge_cells("B2:I2")
    ws["B2"] = "送 货 单"

    ws.merge_cells("B3:I3")
    ws["B3"] = "收货单位：示例收货单位    联系人：示例联系人        NO:SAMPLE-0000"
    ws.merge_cells("B4:I4")
    ws["B4"] = "收货地址：示例收货地址"
    ws.merge_cells("B5:I5")
    ws["B5"] = "联系电话:00000000000                             日期：示例日期"

    headers = ["序号", "商品名称", "型号规格", "数量", "单位", "单价/元", "金额/元", "备注"]
    for col_idx, text in enumerate(headers, start=2):  # starts at column B
        col = openpyxl.utils.get_column_letter(col_idx)
        ws[f"{col}6"] = text

    for row in range(DELIVERY_ITEM_START_ROW, DELIVERY_ITEM_END_ROW + 1):
        ws[f"B{row}"] = row - DELIVERY_ITEM_START_ROW + 1
        ws[f"C{row}"] = "示例产品"
        ws[f"D{row}"] = f"示例型号-{row}"
        ws[f"E{row}"] = 1
        ws[f"F{row}"] = "件"
        ws[f"G{row}"] = f"=H{row}/E{row}"
        ws[f"H{row}"] = 1

    ws.merge_cells("B19:C19")
    ws["B19"] = "合计金额："
    ws.merge_cells("D19:I19")
    ws["D19"] = 999  # stale literal total left by a previous transaction

    ws.merge_cells("B20:I20")

    ws.merge_cells("B21:I21")
    ws["B21"] = "说明：以上货品请核对数量，如有质量问题，请在3日内联系本公司，逾期恕不负责。"

    ws.merge_cells("B22:I22")
    ws["B22"] = "送货单位：示例供应商有限公司"
    ws.merge_cells("B23:I23")
    ws["B23"] = "地址：示例发货地址"
    ws.merge_cells("B24:I24")
    ws["B24"] = "联系电话："

    ws.merge_cells("B25:I25")
    ws["B25"] = "送货单位：示例供应商有限公司         收货单位：示例收货单位 "

    _configure_print(ws, "B1:I25", "1:6", 19)

    wb.save(path)
    return path
