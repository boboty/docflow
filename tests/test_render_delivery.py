from decimal import Decimal
from pathlib import Path

import openpyxl

from docflow.domain.document import DocumentType, build_projection
from docflow.renderers.xlsx import render
from docflow.templates.registry import TemplateRegistry
from tests.conftest import make_fact_pack, make_line_item


def test_single_delivery_note_generation(delivery_template_path: Path, tmp_path: Path):
    pack = make_fact_pack(item_count=3)
    projection = build_projection(pack, DocumentType.DELIVERY_NOTE_V1)
    definition = TemplateRegistry().get(DocumentType.DELIVERY_NOTE_V1)
    output_path = tmp_path / "delivery-note.xlsx"

    render(definition, delivery_template_path, projection, output_path)

    assert output_path.exists()
    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]
    assert ws["B3"].value == (
        f"收货单位：{pack.ship_to.company}    联系人：{pack.ship_to.contact}        NO:{pack.delivery_no}"
    )
    assert ws["B4"].value == f"收货地址：{pack.ship_to.address}"

    first_item = projection.items[0]
    assert ws["B7"].value == 1
    assert ws["C7"].value == first_item.product_name
    assert ws["G7"].value == float(first_item.gross_unit_price)
    # Column H (金额) holds the literal, authoritative gross_amount source
    # fact - never a recomputed =Dn*Fn formula (see mapping comment / P0
    # repair-2 blocker 1).
    assert ws["H7"].value == float(first_item.gross_amount)
    assert ws["D19"].value == float(projection.totals.gross_total)


def test_unused_rows_are_fully_cleared(delivery_template_path: Path, tmp_path: Path):
    """Phase 0 Repair section 3: leftover template rows must not keep any
    stale value/formula that would recalculate to a spurious number.
    """
    pack = make_fact_pack(item_count=3)
    projection = build_projection(pack, DocumentType.DELIVERY_NOTE_V1)
    definition = TemplateRegistry().get(DocumentType.DELIVERY_NOTE_V1)
    output_path = tmp_path / "delivery-note.xlsx"

    render(definition, delivery_template_path, projection, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]

    for offset, item in enumerate(projection.items):
        row = 7 + offset
        assert ws[f"H{row}"].value == float(item.gross_amount)

    for row in range(10, 19):
        for column in "BCDEFGHI":
            assert ws[f"{column}{row}"].value is None


def test_gross_amount_is_written_verbatim_even_when_it_diverges_from_unit_price_times_quantity(
    delivery_template_path: Path, tmp_path: Path
):
    """The P0 regression this guards: when gross_unit_price*quantity and
    gross_amount differ within the documented validation tolerance, the
    rendered delivery note must show the SAME number as the contract would
    (the true gross_amount), not a recomputed D*F that silently disagrees.
    """
    item = make_line_item(1, quantity=42, gross_unit_price="49", gross_amount="2058.10")
    pack = make_fact_pack(item_count=0)
    from dataclasses import replace

    pack = replace(pack, items=(item,))
    projection = build_projection(pack, DocumentType.DELIVERY_NOTE_V1)
    definition = TemplateRegistry().get(DocumentType.DELIVERY_NOTE_V1)
    output_path = tmp_path / "delivery-note.xlsx"

    render(definition, delivery_template_path, projection, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["Sheet1"]
    assert ws["H7"].value == 2058.10
    assert ws["H7"].value != 42 * 49  # the recomputed-formula value this bug used to show
    assert projection.items[0].gross_amount == Decimal("2058.10")
