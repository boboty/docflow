"""XLSX renderer: fills a copy of a real template from a DocumentProjection.

The renderer performs no business calculation. It only:
  - substitutes header text placeholders (document-identity cells),
  - substitutes body-text placeholders (mostly-static clauses that embed
    a transaction fact, e.g. a delivery-deadline sentence naming the
    actual delivery date),
  - writes already-derived item values into fixed template rows,
  - clears leftover sample rows when fewer items than template capacity,
  - fails explicitly when items exceed template capacity.

A template cell must never be left holding a *previous* business
transaction's value: if a cell's content depends on this transaction's
facts at all - even a single date embedded inside an otherwise-static
paragraph - it must be listed in the mapping's `header` or `text` section
so this renderer overwrites it. "Static template boilerplate we don't
touch" only covers cells whose content is genuinely transaction
-independent (legal clause headings, fixed disclaimer wording, etc).

Every written amount is a value already decided upstream (rules.money) -
this renderer never writes a per-row formula that could recompute a
*different* number than the source fact. A real template's own native,
fixed-range aggregate formulas (e.g. contract row 27's
``=SUM(G9:G26)``, delivery's ``D25 = SUM(G7:G24)``) are left untouched;
`fullCalcOnLoad` is set so Excel recalculates them the moment the
generated file is opened.
"""
from __future__ import annotations

import copy
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.workbook.properties import CalcProperties

from docflow.domain.document import DocumentProjection, chinese_date, chinese_month_day
from docflow.rules.money import LineItemAmounts
from docflow.templates.definition import PrintProfile, TemplateDefinition

TEMPLATE_ITEM_CAPACITY_EXCEEDED = "TEMPLATE_ITEM_CAPACITY_EXCEEDED"

# Single source of truth for the placeholder names available to BOTH
# `header` and `text` mapping cells - shared between the actual renderer
# (_cell_text_context) and preflight (which must reject a mapping
# referencing an unknown placeholder *before* any record is processed,
# rather than let it surface as a KeyError mid-batch).
_CELL_TEXT_CONTEXT_KEYS = (
    "buyer", "seller", "contract_no", "delivery_no", "contract_date", "delivery_date",
    "delivery_month_day", "amount_in_words", "ship_to_company", "ship_to_contact",
    "ship_to_phone", "ship_to_address", "seller_contact", "seller_phone", "seller_address",
)

# Field names a mapping's items.columns may legally reference: every
# LineItemAmounts attribute, plus the two renderer-synthesized ones.
_VALID_ITEM_FIELDS = frozenset(
    {"index", "remarks", "sku", "product_name", "specification", "quantity",
     "unit", "gross_unit_price", "net_unit_price", "net_amount", "tax_amount", "gross_amount"}
)


class TemplateRenderError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


class TemplatePreflightError(Exception):
    """A whole-batch-blocking error: the template file or mapping itself is
    unusable. This is distinct from a per-record TemplateRenderError - it
    means no record in the batch could possibly succeed, so batch
    processing must not even start the record loop.
    """

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _is_non_writable_merged_cell(ws, cell_address: str) -> bool:
    """True if writing to this address would hit an openpyxl MergedCell -
    i.e. the address falls inside a merged range but is not that range's
    top-left anchor cell (the only cell in a merge that's actually
    writable).
    """
    row, col = coordinate_to_tuple(cell_address)
    for merged_range in ws.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
            if (row, col) != (merged_range.min_row, merged_range.min_col):
                return True
    return False


def _preflight_cell_text_mapping(ws, mapping: dict[str, str], section_name: str, dummy_context: dict[str, str]) -> None:
    """Shared preflight check for both `header` and `text`: every
    placeholder must be satisfiable and every target cell must be
    writable. Used identically for both sections so a mapping author gets
    the same guarantees regardless of which one a cell lives in.
    """
    for cell_address, template_str in mapping.items():
        try:
            template_str.format(**dummy_context)
        except (KeyError, IndexError, ValueError) as exc:
            raise TemplatePreflightError(
                "TEMPLATE_MAPPING_INVALID",
                f"{section_name} template for {cell_address} is invalid: {exc}",
            ) from exc
        if _is_non_writable_merged_cell(ws, cell_address):
            raise TemplatePreflightError(
                "TEMPLATE_MAPPING_INVALID",
                f"{section_name} cell {cell_address} is inside a merged range but is not its "
                f"top-left cell, so it cannot be written to",
            )


def preflight(definition: TemplateDefinition, template_path: Path) -> None:
    """Verify a template file AND mapping are usable before any record is
    processed. This must catch everything render() could later choke on
    for reasons independent of any specific record's data - an unusable
    template must fail the whole batch up front, not surface mid-loop
    after some records already produced output.

    Deliberately catches broad exceptions from openpyxl when opening the
    file: a corrupt/invalid xlsx can surface as any of several unrelated
    exception types (bad zip, bad XML, missing parts). This is a
    file-usability gate, not a place where hiding a docflow programmer
    error would be a concern.
    """
    if not template_path.exists():
        raise TemplatePreflightError("TEMPLATE_FILE_MISSING", f"template file not found: {template_path}")

    try:
        # Not read_only: merged-cell writability (below) needs
        # ws.merged_cells.ranges, which read_only worksheets don't expose.
        wb = openpyxl.load_workbook(template_path)
    except Exception as exc:
        raise TemplatePreflightError(
            "TEMPLATE_FILE_INVALID", f"cannot open {template_path} as an xlsx workbook: {exc}"
        ) from exc

    if definition.sheet not in wb.sheetnames:
        raise TemplatePreflightError(
            "TEMPLATE_SHEET_MISSING",
            f"sheet {definition.sheet!r} not found in {template_path} (available: {wb.sheetnames})",
        )
    ws = wb[definition.sheet]

    dummy_context = {key: "" for key in _CELL_TEXT_CONTEXT_KEYS}
    _preflight_cell_text_mapping(ws, definition.header, "header", dummy_context)
    _preflight_cell_text_mapping(ws, definition.text, "text", dummy_context)

    unknown_fields = set(definition.items.columns) - _VALID_ITEM_FIELDS
    if unknown_fields:
        raise TemplatePreflightError(
            "TEMPLATE_MAPPING_INVALID",
            f"items.columns references unknown field(s): {sorted(unknown_fields)}",
        )

    for row in range(definition.items.start_row, definition.items.end_row + 1):
        for column_letter in definition.items.columns.values():
            cell_address = f"{column_letter}{row}"
            if _is_non_writable_merged_cell(ws, cell_address):
                raise TemplatePreflightError(
                    "TEMPLATE_MAPPING_INVALID",
                    f"item cell {cell_address} is inside a merged range but is not its "
                    f"top-left cell, so it cannot be written to",
                )

    for cell_address in definition.totals:
        if _is_non_writable_merged_cell(ws, cell_address):
            raise TemplatePreflightError(
                "TEMPLATE_MAPPING_INVALID",
                f"totals cell {cell_address} is inside a merged range but is not its "
                f"top-left cell, so it cannot be written to",
            )


def _cell_text_context(projection: DocumentProjection) -> dict[str, str]:
    return {
        "buyer": projection.buyer,
        "seller": projection.seller,
        "contract_no": projection.contract_no,
        "delivery_no": projection.delivery_no,
        "contract_date": chinese_date(projection.contract_date),
        "delivery_date": chinese_date(projection.delivery_date),
        "delivery_month_day": chinese_month_day(projection.delivery_date),
        "amount_in_words": projection.amount_in_words,
        "ship_to_company": projection.ship_to.company,
        "ship_to_contact": projection.ship_to.contact,
        "ship_to_phone": projection.ship_to.phone,
        "ship_to_address": projection.ship_to.address,
        "seller_contact": projection.seller_contact or "",
        "seller_phone": projection.seller_phone or "",
        "seller_address": projection.seller_address or "",
    }


_NUMERIC_ITEM_FIELDS = {
    "quantity",
    "net_unit_price",
    "gross_unit_price",
    "net_amount",
    "tax_amount",
    "gross_amount",
}


def _item_cell_value(item: LineItemAmounts, index: int, field_name: str):
    if field_name == "index":
        return index
    if field_name == "remarks":
        return item.remarks or ""
    value = getattr(item, field_name)
    if field_name in _NUMERIC_ITEM_FIELDS:
        return float(value)
    return value


def _apply_print_profile(ws, profile: PrintProfile) -> None:
    """Author the sheet's scale-critical print settings from the template's
    own calibrated PrintProfile, rather than trusting whatever the source
    template file happens to carry (typically a user's manual, dynamic "Fit
    to Page" setting - exactly what makes printed seal size depend on
    however that print run's fit-to-page math happened to land). Only the
    settings that actually cause the fit-to-page distortion are touched
    here; everything else (margins, headers/footers, gridlines, breaks) is
    preserved as-is by `_snapshot_print_settings`/`_restore_print_settings`.
    """
    # profile.paper_size is validated at mapping-load time to always be
    # "A4" (see templates.definition._SUPPORTED_PAPER_SIZES) - there is
    # exactly one real-world paper size this project's templates need.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = profile.orientation
    ws.page_setup.scale = profile.scale_percent
    ws.page_setup.fitToWidth = None
    ws.page_setup.fitToHeight = None
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.print_area = profile.print_area


def _snapshot_print_settings(ws) -> dict:
    """Capture every openpyxl-exposed print setting before cell/image edits."""
    return {
        "page_setup": copy.copy(ws.page_setup),
        "page_margins": copy.copy(ws.page_margins),
        "print_options": copy.copy(ws.print_options),
        "page_setup_properties": copy.copy(ws.sheet_properties.pageSetUpPr),
        "print_area": ws.print_area,
        "print_title_rows": ws.print_title_rows,
        "print_title_cols": ws.print_title_cols,
        "row_breaks": copy.deepcopy(ws.row_breaks),
        "col_breaks": copy.deepcopy(ws.col_breaks),
        "odd_header": copy.copy(ws.oddHeader),
        "odd_footer": copy.copy(ws.oddFooter),
        "even_header": copy.copy(ws.evenHeader),
        "even_footer": copy.copy(ws.evenFooter),
        "first_header": copy.copy(ws.firstHeader),
        "first_footer": copy.copy(ws.firstFooter),
    }


def _restore_print_settings(ws, settings: dict) -> None:
    ws.page_setup = settings["page_setup"]
    ws.page_margins = settings["page_margins"]
    ws.print_options = settings["print_options"]
    ws.sheet_properties.pageSetUpPr = settings["page_setup_properties"]
    ws.print_area = settings["print_area"]
    ws.print_title_rows = settings["print_title_rows"]
    ws.print_title_cols = settings["print_title_cols"]
    ws.row_breaks = settings["row_breaks"]
    ws.col_breaks = settings["col_breaks"]
    ws.oddHeader = settings["odd_header"]
    ws.oddFooter = settings["odd_footer"]
    ws.evenHeader = settings["even_header"]
    ws.evenFooter = settings["even_footer"]
    ws.firstHeader = settings["first_header"]
    ws.firstFooter = settings["first_footer"]


def _insert_images(ws, definition: TemplateDefinition, image_assets: dict[str, Path]) -> tuple[str, ...]:
    # A mapping's printed_diameter_mm is the seal's target size on the
    # PRINTED page. Since the whole sheet is printed at print.scale_percent,
    # the image must be embedded into the workbook larger than that by the
    # inverse of the scale factor, so it still measures printed_diameter_mm
    # once the page itself is scaled down (e.g. at 70% scale, a 38mm printed
    # target is embedded at 38 / 0.70 ≈ 54.29mm in the workbook - looking
    # oversized in Excel is expected and necessary, not a bug).
    scale_factor = definition.print_profile.scale_percent / 100.0
    enhancements: list[str] = []
    for image_id, mapping in definition.images.items():
        asset_path = image_assets.get(image_id)
        if asset_path is None:
            enhancements.append(f"IMAGE_SKIPPED {image_id}: asset not configured")
            continue
        try:
            image = Image(asset_path)
            column_letters, row = openpyxl.utils.cell.coordinate_from_string(mapping.anchor)
            marker = AnchorMarker(
                col=column_index_from_string(column_letters) - 1,
                row=row - 1,
                colOff=pixels_to_EMU(mapping.x_offset_px),
                rowOff=pixels_to_EMU(mapping.y_offset_px),
            )
            workbook_diameter_mm = mapping.printed_diameter_mm / scale_factor
            size_emu = round(workbook_diameter_mm * 36_000)
            image.anchor = OneCellAnchor(
                _from=marker,
                ext=XDRPositiveSize2D(cx=size_emu, cy=size_emu),
            )
            ws.add_image(image)
            enhancements.append(f"IMAGE_INSERTED {image_id}")
        except Exception as exc:
            # Images are a delivery enhancement. A stale/malformed asset
            # must never turn an otherwise valid business document into FAIL.
            enhancements.append(f"IMAGE_SKIPPED {image_id}: {exc}")
    return tuple(enhancements)


def render(
    definition: TemplateDefinition,
    source_template_path: Path,
    projection: DocumentProjection,
    output_path: Path,
    image_assets: dict[str, Path] | None = None,
) -> tuple[str, ...]:
    capacity = definition.items.capacity
    if len(projection.items) > capacity:
        raise TemplateRenderError(
            TEMPLATE_ITEM_CAPACITY_EXCEEDED,
            f"{len(projection.items)} line items exceed template capacity {capacity} "
            f"for {definition.id}",
        )

    wb = openpyxl.load_workbook(source_template_path)
    ws = wb[definition.sheet]
    print_settings = _snapshot_print_settings(ws)

    context = _cell_text_context(projection)
    for cell_address, template_str in definition.header.items():
        ws[cell_address] = template_str.format(**context)
    for cell_address, template_str in definition.text.items():
        ws[cell_address] = template_str.format(**context)

    columns = definition.items.columns
    start_row = definition.items.start_row
    for offset, item in enumerate(projection.items):
        row = start_row + offset
        for field_name, column_letter in columns.items():
            ws[f"{column_letter}{row}"] = _item_cell_value(item, offset + 1, field_name)

    for row in range(start_row + len(projection.items), definition.items.end_row + 1):
        for column_letter in columns.values():
            ws[f"{column_letter}{row}"] = None

    for cell_address, field_name in definition.totals.items():
        ws[cell_address] = float(getattr(projection.totals, field_name))

    enhancements = _insert_images(ws, definition, image_assets or {})
    # Reapply the source template's settings after all mutations (margins,
    # headers/footers, gridlines, breaks - a copy, not a guess). Then
    # authoritatively overwrite the scale-critical subset from the
    # template's own calibrated PrintProfile: fixed A4/portrait/print
    # area/scale, fitToPage disabled - never whatever "Fit to Page" state
    # the source file happened to carry.
    _restore_print_settings(ws, print_settings)
    _apply_print_profile(ws, definition.print_profile)

    if wb.calculation is None:
        wb.calculation = CalcProperties(fullCalcOnLoad=True)
    else:
        wb.calculation.fullCalcOnLoad = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return enhancements
